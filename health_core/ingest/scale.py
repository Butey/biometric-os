"""Импорт выгрузок весов Feelfit (.xlsx/.csv) в body_metrics — §04, §08, §12.

Две ловушки:
  1. «нуля»: 0 и прочерк оба значат «не снято» -> NULL, кроме visceral_fat (0 там валиден).
  2. «серий»: замеры кластеризуются по разрыву <= burst_gap_sec, не по усечению до минуты.
"""
import csv
import hashlib
import statistics
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

from health_core.config import load

# Столбец в выгрузке -> поле body_metrics (§04). Сопоставление дословное, опечатку
# «мыщцы» в исходнике исправлять нельзя. «Имя устройства» в схему не входит — игнорируем.
COLUMNS = {
    "Время измерения": "measured_at",
    "Вес(kg)": "weight_kg",
    "Содержание жира(%)": "fat_pct",
    "Индекс массы тела": "bmi",
    "Скелетные мыщцы(%)": "skeletal_muscle_pct",
    "Мышечная масса(kg)": "muscle_mass_kg",
    "Белки(%)": "protein_pct",
    "Скорость обмена веществ(kcal)": "device_bmr_kcal",
    "Масса тела без учета жира(kg)": "ffm_kg",
    "Подкожно-жировая клетчатка(%)": "subcutaneous_fat_pct",
    "Висцеральный жир": "visceral_fat",
    "Содержание воды в организме(%)": "water_pct",
    "Костная масса(kg)": "bone_mass_kg",
    "Метаболический возраст": "metabolic_age",
    "MAC-адрес устройства": "device_mac",
}
NUMERIC_FIELDS = [v for v in COLUMNS.values() if v not in ("measured_at", "device_mac")]


def _sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _to_float(raw) -> float | None:
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "-"):
        return None
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return None


def _read_rows(path: str):
    """Отдаёт по строке сырых значений, ключи — канонические имена полей."""
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            header = next(reader)
            idx = {name: i for i, name in enumerate(header)}
            for row in reader:
                yield {
                    field: row[idx[col]] if idx[col] < len(row) else None
                    for col, field in COLUMNS.items() if col in idx
                }
    else:
        # openpyxl отвергает файлы по расширению .xls, даже если внутри настоящий xlsx-zip
        # (как здесь — реальная выгрузка называется "*.xlsx.xls"). Обход: отдаём файловый
        # объект вместо строки-пути, тогда проверка расширения в openpyxl не срабатывает.
        with open(path, "rb") as fh:
            wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
            try:
                ws = wb[wb.sheetnames[0]]
                rows = ws.iter_rows(values_only=True)
                header = next(rows)
                idx = {name: i for i, name in enumerate(header) if name is not None}
                for row in rows:
                    yield {
                        field: row[idx[col]] if idx[col] < len(row) else None
                        for col, field in COLUMNS.items() if col in idx
                    }
            finally:
                wb.close()


def _parse_row(raw: dict) -> dict | None:
    """Строка без времени или веса отбрасывается. Ноль -> NULL везде, кроме visceral_fat."""
    dt_raw = raw.get("measured_at")
    if not dt_raw:
        return None
    try:
        measured_at = datetime.strptime(str(dt_raw).strip(), "%d/%m/%Y %H:%M:%S")
    except ValueError:
        return None

    weight_kg = _to_float(raw.get("weight_kg"))
    if not weight_kg:
        return None

    values = {"measured_at": measured_at, "weight_kg": weight_kg}
    for field in NUMERIC_FIELDS:
        if field == "weight_kg":
            continue
        v = _to_float(raw.get(field))
        if field != "visceral_fat" and v == 0:
            v = None
        values[field] = v
    mac = raw.get("device_mac")
    values["device_mac"] = str(mac).strip() if mac not in (None, "") else None
    return values


def _cluster_bursts(rows: list[dict], gap_sec: int) -> list[dict]:
    """Серия := замеры подряд с разрывом <= gap_sec к предыдущему (не усечение до минуты)."""
    rows = sorted(rows, key=lambda r: r["measured_at"])
    bursts: list[list[dict]] = []
    for row in rows:
        if bursts and (row["measured_at"] - bursts[-1][-1]["measured_at"]) <= timedelta(seconds=gap_sec):
            bursts[-1].append(row)
        else:
            bursts.append([row])

    collapsed = []
    for burst in bursts:
        first_ts = burst[0]["measured_at"].strftime("%Y-%m-%d %H:%M:%S")
        out = {"burst_key": first_ts, "measured_at": first_ts}
        for field in NUMERIC_FIELDS:
            vals = [r[field] for r in burst if r[field] is not None]
            out[field] = statistics.median(vals) if vals else None
        out["device_mac"] = next((r["device_mac"] for r in burst if r["device_mac"]), None)
        collapsed.append(out)
    return collapsed


def import_export(conn, user_id: int, path: str) -> dict:
    """Импорт одной выгрузки. Возвращает {"added": n, "skipped": n, "bursts": n}."""
    file_hash = _sha256(path)
    if conn.execute("SELECT 1 FROM import_log WHERE file_hash=?", (file_hash,)).fetchone():
        return {"added": 0, "skipped": 0, "bursts": 0}

    gap_sec = load().get("ingest", {}).get("burst_gap_sec", 120)

    rows = []
    for raw in _read_rows(path):
        parsed = _parse_row(raw)
        if parsed is not None:
            rows.append(parsed)
    bursts = _cluster_bursts(rows, gap_sec)

    other_fields = [f for f in NUMERIC_FIELDS if f != "weight_kg"]
    cols = ["user_id", "burst_key", "measured_at", "weight_kg", *other_fields, "device_mac"]
    sql = (
        f"INSERT INTO body_metrics({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))}) "
        "ON CONFLICT(user_id, burst_key) DO NOTHING"
    )

    added = skipped = 0
    for b in bursts:
        values = [user_id, b["burst_key"], b["measured_at"], b["weight_kg"]]
        values += [b[f] for f in other_fields]
        values.append(b["device_mac"])
        cur = conn.execute(sql, values)
        if cur.rowcount:
            added += 1
        else:
            skipped += 1

    if len(bursts) > 0:
        conn.execute(
            "INSERT INTO import_log(user_id, imported_at, file_hash) VALUES (?, ?, ?) "
            "ON CONFLICT(file_hash) DO NOTHING",
            (user_id, datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"), file_hash),
        )
    conn.commit()
    return {"added": added, "skipped": skipped, "bursts": len(bursts)}


if __name__ == "__main__":
    import csv as _csv
    import os
    import tempfile

    def raw(measured_at, weight="80.0", fat="20.0", visceral="10", mac="AA:BB"):
        return {
            "measured_at": measured_at, "weight_kg": weight, "fat_pct": fat,
            "bmi": "25.0", "skeletal_muscle_pct": "40.0", "muscle_mass_kg": "76.0",
            "protein_pct": "14.0", "device_bmr_kcal": "2000", "ffm_kg": "80.0",
            "subcutaneous_fat_pct": "30.0", "visceral_fat": visceral, "water_pct": "45.0",
            "bone_mass_kg": "4.0", "metabolic_age": "40", "device_mac": mac,
        }

    # --- часть 1: кластеризация в памяти, без файлов и БД ---
    straddle = [
        raw("20/08/2026 08:31:58", weight="80.5"),
        raw("20/08/2026 08:32:03", weight="80.3"),
        raw("20/08/2026 08:32:05", weight="80.7"),
    ]
    later = raw("20/08/2026 08:42:10", weight="79.0", fat="0", visceral="0")  # проверка ловушки нуля

    parsed = [p for p in (_parse_row(r) for r in straddle + [later]) if p is not None]
    bursts = _cluster_bursts(parsed, gap_sec=120)

    assert len(bursts) == 2, f"expected 2 bursts (straddle collapsed + later separate), got {len(bursts)}"
    b0, b1 = bursts
    assert b0["burst_key"] == "2026-08-20 08:31:58", b0["burst_key"]  # ключ = первый замер серии
    assert b0["weight_kg"] == statistics.median([80.5, 80.3, 80.7]) == 80.5
    assert b1["burst_key"] == "2026-08-20 08:42:10"
    assert b1["fat_pct"] is None, "0 в fat_pct должен стать NULL"
    assert b1["visceral_fat"] == 0, "0 в visceral_fat — валидное значение, должно сохраниться как 0"
    print("OK part1: разрыв-кластеризация (не минутная), медианы, ловушка нуля")

    # --- часть 2: реальный import_export через временную БД, дважды подряд ---
    with tempfile.TemporaryDirectory() as tmp:
        os.environ["HEALTH_DB"] = str(Path(tmp) / "health.db")
        from health_core import db
        db.DB_PATH = Path(os.environ["HEALTH_DB"])
        conn = db.connect()
        db.migrate(conn)
        conn.execute(
            "INSERT INTO users(telegram_user_id, created_at) VALUES (1, '2026-08-20 00:00:00')"
        )
        conn.commit()
        uid = conn.execute("SELECT id FROM users WHERE telegram_user_id=1").fetchone()["id"]

        csv_path = Path(tmp) / "export.csv"
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            w = _csv.DictWriter(f, fieldnames=list(COLUMNS.keys()))
            w.writeheader()
            for r in (straddle + [later]):
                w.writerow({col: r[field] for col, field in COLUMNS.items()})

        first = import_export(conn, uid, str(csv_path))
        assert first == {"added": 2, "skipped": 0, "bursts": 2}, first
        second = import_export(conn, uid, str(csv_path))
        assert second == {"added": 0, "skipped": 0, "bursts": 0}, second
        total = conn.execute("SELECT COUNT(*) c FROM body_metrics WHERE user_id=?", (uid,)).fetchone()["c"]
        assert total == 2, f"expected 2 rows after two imports of the same file, got {total}"
        conn.close()
        print(f"OK part2: import_export first={first} second={second} total_rows={total}")

    print("ALL OK")
